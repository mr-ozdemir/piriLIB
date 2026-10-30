"""Plans as CSV (and Excel) tables, with names instead of indices.

One row per load. Rows of the same dispatch repeat its details; a dispatch
without loads has one row with the load columns empty. Columns:

``dispatch``
    Dispatch number, from 1.
``date``, ``kind``, ``vehicle_type``, ``origin``
    The day of the dispatch, ``owned`` or ``spot``, the vehicle type and the origin.
``start``
    When loading starts, ``HH:MM`` after 00:00 of ``date`` (hours may pass 24).
    Optional when reading; empty means as soon as every load is ready.
``source``
    The fleet entry, ``fleet.owned[0]`` or ``fleet.spot[1]`` as in the problem file.
``stops``
    The stops in visiting order, separated by ``|``.
``route_km``, ``dispatch_cost``
    Filled in by :func:`write_plan`; ignored by :func:`read_plan`.
``demand``, ``destination``, then one column per quantity dimension
    The load.
"""

import csv
import datetime
import os
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any, Literal

import numpy as np

from pirilib.errors import PlanError
from pirilib.evaluation import Evaluation
from pirilib.plan import SPOT, Dispatch, Kind, Load, Plan
from pirilib.spec._format import bullets
from pirilib.spec.instance import LineHaulInstance
from pirilib.spec.tables import LIST_SEPARATOR

__all__ = [
    "columns",
    "read_plan",
    "rows",
    "summary",
    "write_plan",
]

_FIXED = ("dispatch", "date", "start", "kind", "vehicle_type", "source", "origin", "stops")
_OPTIONAL = ("start",)
_CLOCK = re.compile(r"(\d+):(\d{2})(?::(\d{2}(?:\.\d+)?))?")
_COMPUTED = ("route_km", "dispatch_cost")
_LOAD = ("demand", "destination")
_SOURCE = re.compile(r"fleet\.(owned|spot)\[(\d+)\]")


def columns(instance: LineHaulInstance) -> tuple[str, ...]:
    """The column names of the plan table, ending with the dimensions of ``instance``."""
    return (*_FIXED, *_COMPUTED, *_LOAD, *instance.dimensions)


def rows(evaluation: Evaluation) -> Iterator[dict[str, Any]]:
    """The rows of the plan table, as values (numbers are floats)."""
    instance, plan = evaluation.instance, evaluation.plan
    for position, dispatch in enumerate(plan.dispatches()):
        kind_types = instance.owned_type if dispatch.kind == "owned" else instance.spot_type
        common: dict[str, Any] = {
            "dispatch": position + 1,
            "date": instance.days[dispatch.day].isoformat(),
            "start": _clock(float(evaluation.start_minute[position]) - 1440.0 * dispatch.day),
            "kind": dispatch.kind,
            "vehicle_type": instance.vehicle_types[int(kind_types[dispatch.source])],
            "source": f"fleet.{dispatch.kind}[{dispatch.source}]",
            "origin": instance.centers[dispatch.origin],
            "stops": LIST_SEPARATOR.join(instance.centers[stop] for stop in dispatch.stops),
            "route_km": float(evaluation.route_km[position]),
            "dispatch_cost": float(evaluation.dispatch_cost[position]),
        }
        if not dispatch.loads:
            yield {**common, **dict.fromkeys((*_LOAD, *instance.dimensions), "")}
        for load in dispatch.loads:
            yield {
                **common,
                "demand": instance.demand_ids[load.demand],
                "destination": instance.centers[int(instance.demand_destination[load.demand])],
                **dict(zip(instance.dimensions, load.quantity, strict=True)),
            }


def _clock(minutes: float) -> str:
    whole, seconds = divmod(round(minutes * 60.0, 6), 60.0)
    hours, rest = divmod(int(whole), 60)
    text = f"{hours:02d}:{rest:02d}"
    return text if seconds == 0 else f"{text}:{seconds:09.6f}".rstrip("0").rstrip(".")


def _minutes_of(text: str) -> float | None:
    match = _CLOCK.fullmatch(text)
    if match is None:
        return None
    seconds = float(match.group(3) or 0.0)
    return int(match.group(1)) * 60.0 + int(match.group(2)) + seconds / 60.0


def write_plan(
    evaluation: Evaluation,
    path: str | os.PathLike[str],
    *,
    delimiter: str = ",",
    decimal: Literal[".", ","] = ".",
) -> None:
    """Write the plan of ``evaluation`` as CSV, or as Excel if ``path`` ends with ``.xlsx``.

    Excel files need openpyxl (``pip install "pirilib[excel]"``).
    """
    file = Path(path)
    header = columns(evaluation.instance)
    if file.suffix.lower() == ".xlsx":
        _write_excel(file, header, rows(evaluation))
        return
    with file.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=header, delimiter=delimiter)
        writer.writeheader()
        for row in rows(evaluation):
            writer.writerow({key: _text(value, decimal) for key, value in row.items()})


def _text(value: Any, decimal: str) -> str:
    if isinstance(value, float):
        text = repr(value)
        return text.replace(".", ",") if decimal == "," else text
    return str(value)


def _write_excel(file: Path, header: tuple[str, ...], data: Iterator[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        raise ImportError(
            'writing .xlsx files needs openpyxl: pip install "pirilib[excel]"'
        ) from None
    book = Workbook()
    sheet = book.active
    sheet.title = "plan"
    sheet.append(list(header))
    for row in data:
        sheet.append([row[name] for name in header])
    book.save(file)


def read_plan(
    instance: LineHaulInstance,
    path: str | os.PathLike[str],
    *,
    delimiter: str = ",",
    decimal: Literal[".", ","] = ".",
) -> Plan:
    """Read a plan written by :func:`write_plan` (or by hand) for ``instance``.

    A file ending with ``.xlsx`` is read as Excel (its ``plan`` sheet, or the
    first one), which needs openpyxl (``pip install "pirilib[excel]"``). Every
    problem is reported at once in a :class:`~pirilib.PlanError`, with the line
    (or Excel row) it is on.
    """
    file = Path(path)
    days = {day.isoformat(): position for position, day in enumerate(instance.days)}
    centers = {center: position for position, center in enumerate(instance.centers)}
    demands = {demand: position for position, demand in enumerate(instance.demand_ids)}
    sizes = {"owned": instance.owned_type.size, "spot": instance.spot_type.size}
    problems: list[str] = []
    order: list[str] = []
    parts: dict[str, dict[str, Any]] = {}
    try:
        if file.suffix.lower() == ".xlsx":
            header, records = _read_excel(file, decimal)
        else:
            header, records = _read_csv(file, delimiter)
    except OSError as exc:
        raise PlanError(f"cannot read {file}: {exc.strerror or exc}") from exc
    except UnicodeDecodeError as exc:
        raise PlanError(f"{file}: not valid UTF-8 text: {exc.reason} at byte {exc.start}") from exc
    except csv.Error as exc:
        raise PlanError(f"{file.name}: invalid CSV: {exc}") from exc
    missing = [
        name
        for name in (*_FIXED, *_LOAD, *instance.dimensions)
        if name not in header and name not in _OPTIONAL
    ]
    if missing:
        hint = " (is the delimiter right?)" if len(header) <= 1 else ""
        raise PlanError(f"{file.name}: missing column(s) {missing}{hint}")
    for line, record in records:
        where = f"{file.name}:{line}"
        found = _read_row(record, instance, days, centers, demands, sizes, decimal)
        if isinstance(found, list):
            problems += [f"{where}: {problem}" for problem in found]
            continue
        key, dispatch, load = found
        if key not in parts:
            order.append(key)
            parts[key] = {"dispatch": dispatch, "loads": []}
        elif parts[key]["dispatch"] != dispatch:
            problems.append(f"{where}: dispatch {key} differs from its first row")
        if load is not None:
            parts[key]["loads"].append(load)
    if problems:
        raise PlanError(f"{file.name}: invalid plan\n{bullets(problems)}")
    dispatches = [
        Dispatch(**{**parts[key]["dispatch"], "loads": tuple(parts[key]["loads"])}) for key in order
    ]
    return Plan.from_dispatches(dispatches, len(instance.dimensions))


Records = list[tuple[int, dict[str, str | None]]]
"""Rows of a plan table with their line (or Excel row) numbers."""


def _read_csv(file: Path, delimiter: str) -> tuple[list[str], Records]:
    with file.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, delimiter=delimiter)
        records = [(reader.line_num, record) for record in reader]
        return list(reader.fieldnames or []), records


def _read_excel(file: Path, decimal: str) -> tuple[list[str], Records]:
    """The ``plan`` sheet (or the first one), every cell as the text a CSV file would hold."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError(
            'reading .xlsx files needs openpyxl: pip install "pirilib[excel]"'
        ) from None
    try:
        book = load_workbook(file, read_only=True, data_only=True)
    except zipfile.BadZipFile as exc:
        raise PlanError(f"cannot read {file}: not an Excel workbook") from exc
    try:
        sheet = book["plan"] if "plan" in book.sheetnames else book.worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if not values:
        return [], []
    header = [_cell_text(value, decimal) for value in values[0]]
    records: Records = [
        (
            line,
            {
                name: _cell_text(value, decimal)
                for name, value in zip(header, row, strict=False)
                if name
            },
        )
        for line, row in enumerate(values[1:], start=2)
        if any(value not in (None, "") for value in row)
    ]
    return header, records


def _cell_text(value: Any, decimal: str) -> str:
    """An Excel value as CSV text: dates as 2026-01-05, times and durations as 09:30."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        return _clock(value.hour * 60.0 + value.minute + value.second / 60.0)
    if isinstance(value, datetime.timedelta):
        return _clock(value.total_seconds() / 60.0)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value, decimal)


def _read_row(
    record: dict[str, str | None],
    instance: LineHaulInstance,
    days: dict[str, int],
    centers: dict[str, int],
    demands: dict[str, int],
    sizes: dict[str, int],
    decimal: str,
) -> tuple[str, dict[str, Any], Load | None] | list[str]:
    cell = {key: (value or "").strip() for key, value in record.items() if key is not None}
    problems: list[str] = []
    key = cell.get("dispatch", "")
    if not key:
        problems.append("'dispatch' is empty")
    start = None
    if cell.get("start"):
        start = _minutes_of(cell["start"])
        if start is None:
            problems.append(f"start '{cell['start']}' is not like 09:30")
    day = days.get(cell.get("date", ""))
    if day is None:
        problems.append(f"date '{cell.get('date', '')}' is not in the calendar")
    match = _SOURCE.fullmatch(cell.get("source", ""))
    kind: Kind = "spot"
    source = 0
    if match is None:
        problems.append(f"source '{cell.get('source', '')}' is not like fleet.spot[0]")
    else:
        kind = "owned" if match.group(1) == "owned" else "spot"
        source = int(match.group(2))
        if source >= sizes[kind]:
            problems.append(f"there is no fleet.{kind}[{source}]")
        else:
            types = instance.owned_type if kind == "owned" else instance.spot_type
            expected = instance.vehicle_types[int(types[source])]
            if cell.get("kind", kind) != kind:
                problems.append(f"kind '{cell['kind']}' does not match source {cell['source']}")
            if cell.get("vehicle_type", expected) != expected:
                problems.append(
                    f"vehicle_type '{cell['vehicle_type']}' does not match source "
                    f"{cell['source']} ({expected})"
                )
    names = [cell.get("origin", "")] + [
        stop.strip() for stop in cell.get("stops", "").split(LIST_SEPARATOR) if stop.strip()
    ]
    unknown = [name for name in names if name not in centers]
    if unknown:
        problems.append(f"unknown center(s) {unknown}")
    load = None
    if cell.get("demand"):
        demand = demands.get(cell["demand"])
        if demand is None:
            problems.append(f"unknown demand '{cell['demand']}'")
        quantities = []
        for name in instance.dimensions:
            text = cell.get(name, "").replace(" ", "")
            if decimal == ",":
                text = text.replace(".", "").replace(",", ".")
            try:
                quantities.append(float(text))
            except ValueError:
                problems.append(f"'{name}': '{cell.get(name, '')}' is not a number")
        if demand is not None and len(quantities) == len(instance.dimensions):
            load = Load(demand, tuple(quantities))
    if problems:
        return problems
    assert day is not None
    dispatch = {
        "day": day,
        "kind": kind,
        "source": source,
        "origin": centers[names[0]],
        "stops": tuple(centers[name] for name in names[1:]),
        "start": start,
    }
    return key, dispatch, load


def summary(evaluation: Evaluation) -> dict[str, Any]:
    """Headline numbers of an evaluation, for reports."""
    plan = evaluation.plan
    return {
        "total_cost": evaluation.total_cost,
        "owned_fixed_cost": evaluation.owned_fixed_cost,
        "service_cost": evaluation.service_cost,
        "dispatches": len(plan),
        "spot_dispatches": int(np.count_nonzero(plan.kind == SPOT)),
        "total_km": evaluation.total_km,
        "unserved": float(evaluation.unserved.sum()),
        "feasible": evaluation.feasible,
        "violations": evaluation.violated(),
    }
