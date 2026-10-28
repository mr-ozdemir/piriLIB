from pathlib import Path
from typing import Any

import numpy as np
import pytest

import pirilib
from pirilib import PlanError, check_plan, evaluate
from pirilib.planio import read_plan, rows, summary, write_plan
from pirilib.spec import LineHaulInstance


@pytest.fixture
def solved(example_path: Path) -> pirilib.SolveResult:
    spec = pirilib.load_spec(example_path)
    return pirilib.solve(spec.compile(), spec.build_objective(), seed=1)


def same_plan(first: pirilib.Plan, second: pirilib.Plan) -> bool:
    return all(
        np.array_equal(getattr(first, name), getattr(second, name))
        for name in (
            "day",
            "kind",
            "source",
            "origin",
            "stop_start",
            "stop_center",
            "load_demand",
            "load_quantity",
        )
    )


@pytest.mark.parametrize(("delimiter", "decimal"), [(",", "."), (";", ",")])
def test_round_trip(
    tmp_path: Path, solved: pirilib.SolveResult, delimiter: str, decimal: str
) -> None:
    instance = solved.evaluation.instance
    path = tmp_path / "plan.csv"
    write_plan(solved.evaluation, path, delimiter=delimiter, decimal=decimal)
    plan = read_plan(instance, path, delimiter=delimiter, decimal=decimal)
    assert same_plan(plan, solved.plan)
    assert check_plan(instance, plan).feasible


def test_table_layout(tmp_path: Path, solved: pirilib.SolveResult) -> None:
    path = tmp_path / "plan.csv"
    write_plan(solved.evaluation, path)
    header, first, *_ = path.read_text(encoding="utf-8-sig").splitlines()
    assert header == (
        "dispatch,date,start,kind,vehicle_type,source,origin,stops,route_km,dispatch_cost,"
        "demand,destination,desi"
    )
    assert first.startswith("1,2026-05-11,")


def test_dispatches_without_loads_are_kept(tmp_path: Path, instance: LineHaulInstance) -> None:
    plan = pirilib.Plan.from_dispatches(
        [pirilib.Dispatch(day=1, kind="spot", source=0, origin=0, stops=(1,))], 1
    )
    path = tmp_path / "plan.csv"
    write_plan(evaluate(instance, plan), path)
    assert same_plan(read_plan(instance, path), plan)


def test_problems_are_reported_with_lines(tmp_path: Path, instance: LineHaulInstance) -> None:
    path = tmp_path / "plan.csv"
    path.write_text(
        "dispatch,date,kind,vehicle_type,source,origin,stops,demand,destination,desi\n"
        "1,2026-01-05,spot,big,fleet.spot[0],A,B,D000001,B,900\n"
        "2,2026-01-09,spot,big,fleet.spot[0],A,X,D000009,B,abc\n"
        "3,2026-01-05,spot,small,fleet.spot[0],A,B,,,\n"
        "4,2026-01-05,owned,big,fleet.owned[3],A,B,,,\n",
        encoding="utf-8",
    )
    with pytest.raises(PlanError) as info:
        read_plan(instance, path)
    assert str(info.value).splitlines() == [
        "plan.csv: invalid plan",
        "- plan.csv:3: date '2026-01-09' is not in the calendar",
        "- plan.csv:3: unknown center(s) ['X']",
        "- plan.csv:3: unknown demand 'D000009'",
        "- plan.csv:3: 'desi': 'abc' is not a number",
        "- plan.csv:4: vehicle_type 'small' does not match source fleet.spot[0] (big)",
        "- plan.csv:5: there is no fleet.owned[3]",
    ]


def test_excel(tmp_path: Path, solved: pirilib.SolveResult) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "plan.xlsx"
    write_plan(solved.evaluation, path)
    sheet = load_workbook(path).active
    assert sheet.title == "plan"
    assert sheet.max_row == 1 + len(list(rows(solved.evaluation)))
    plan = read_plan(solved.evaluation.instance, path)
    assert same_plan(plan, solved.plan)


HAND_COLUMNS = [
    "kind",
    "vehicle_type",
    "source",
    "origin",
    "stops",
    "demand",
    "destination",
    "desi",
]


def test_excel_edited_by_hand(tmp_path: Path, instance: LineHaulInstance) -> None:
    """Excel turns typed dates and times into its own values; they read as in CSV."""
    import datetime

    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "notes"
    sheet = book.create_sheet("plan")
    sheet.append(["dispatch", "date", "start", *HAND_COLUMNS])
    sheet.append(
        [
            1,
            datetime.datetime(2026, 1, 5),
            datetime.time(9, 30),
            "spot",
            "big",
            "fleet.spot[0]",
            "A",
            "B",
            "D000001",
            "B",
            900,
        ]
    )
    sheet.append([])
    sheet.append(
        [
            2,
            datetime.date(2026, 1, 5),
            datetime.timedelta(hours=26, minutes=30),
            "spot",
            "small",
            "fleet.spot[1]",
            "A",
            "C",
            "D000002",
            "C",
            250.5,
        ]
    )
    path = tmp_path / "plan.xlsx"
    book.save(path)
    plan = read_plan(instance, path)
    assert plan.day.tolist() == [0, 0]
    assert plan.start_minutes.tolist() == [570.0, 1590.0]
    assert plan.load_quantity[:, 0].tolist() == [900.0, 250.5]


def test_excel_problems_name_the_row(tmp_path: Path, instance: LineHaulInstance) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(["dispatch", "date", *HAND_COLUMNS])
    book.active.append(
        [1, "2026-01-05", "spot", "big", "fleet.spot[9]", "A", "B", "D000001", "B", 9]
    )
    path = tmp_path / "plan.xlsx"
    book.save(path)
    with pytest.raises(PlanError, match=r"plan\.xlsx:2: there is no fleet\.spot\[9\]"):
        read_plan(instance, path)
    path.write_text("dispatch,date\n", encoding="utf-8")
    with pytest.raises(PlanError, match="not an Excel workbook"):
        read_plan(instance, path)


def test_summary(solved: pirilib.SolveResult) -> None:
    numbers = summary(solved.evaluation)
    assert numbers["total_cost"] == pytest.approx(solved.evaluation.total_cost)
    assert numbers["dispatches"] == len(solved.plan)
    assert numbers["feasible"] is True
    assert numbers["violations"] == {}


def test_start_times_round_trip(tmp_path: Path, document: dict, compile_document: Any) -> None:
    document["schema"] = "pirilib/linehaul@0.2"
    document["vehicle_types"][0]["speed_kmh"] = 60.0
    document["vehicle_types"][1]["speed_kmh"] = 60.0
    document["fleet"]["owned"][0]["cost"]["per_hour"] = 10.0
    document["demand"][0]["ready"] = "09:00"
    document["rules"] = {"service": {"due_hours": 3, "penalty": 0.4, "max_deferral_days": 1}}
    instance = compile_document(document)
    plan = pirilib.Plan.from_dispatches(
        [
            pirilib.Dispatch(0, "owned", 0, 0, (1,), (pirilib.Load(0, (900.0,)),), start=600.0),
            pirilib.Dispatch(0, "spot", 1, 0, (2,), (pirilib.Load(1, (250.0,)),), start=1590.0),
            pirilib.Dispatch(1, "spot", 1, 1, (0,), (pirilib.Load(2, (120.0,)),)),
        ],
        1,
    )
    first = evaluate(instance, plan)
    path = tmp_path / "plan.csv"
    write_plan(first, path)
    _, first_row, second_row, third_row = path.read_text("utf-8-sig").splitlines()
    starts = [first_row.split(",")[2], second_row.split(",")[2], third_row.split(",")[2]]
    assert starts == ["10:00", "26:30", "00:00"]
    again = evaluate(instance, read_plan(instance, path))
    assert again.total_cost == pytest.approx(first.total_cost)
    assert again.load_late_hours.tolist() == first.load_late_hours.tolist()
    assert again.violated() == first.violated()


def test_start_column_is_optional(tmp_path: Path, instance: LineHaulInstance) -> None:
    path = tmp_path / "plan.csv"
    path.write_text(
        "dispatch,date,kind,vehicle_type,source,origin,stops,demand,destination,desi\n"
        "1,2026-01-05,spot,big,fleet.spot[0],A,B,D000001,B,900\n",
        encoding="utf-8",
    )
    assert np.isnan(read_plan(instance, path).start_minutes).all()
